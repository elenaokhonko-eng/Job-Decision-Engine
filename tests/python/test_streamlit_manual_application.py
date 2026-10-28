import base64
import hashlib
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import docx
import pytest
from streamlit.testing.v1 import AppTest

from streamlit_manual_application import (
    _download_bytes,
    _validate_payload,
    extract_uploaded_description,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]


def test_manual_application_requires_comparison_then_explicit_approval_and_downloads():
    app = AppTest.from_file(
        str(PROJECT_ROOT / "tests/python/manual_application_fixture.py")
    ).run(timeout=10)
    assert not app.exception, [(error.type, error.message) for error in app.exception]

    app.text_input(key="manual_job_url").set_value("https://example.test/job")
    app.text_input(key="manual_job_title").set_value("Data Architect")
    app.text_input(key="manual_job_company").set_value("Example")
    app.button(key="manual_fetch_description").click().run(timeout=10)
    assert (
        "fetched from the selected public posting"
        in app.text_area(key="manual_job_description").value
    )
    assert app.session_state.get("manual_job_source_kind") == "FETCHED"
    assert not any(
        button.label == "Generate customised CV and cover letter"
        for button in app.button
    )

    app.button(key="manual_compare").click().run(timeout=10)
    assert not app.exception, [(error.type, error.message) for error in app.exception]
    assert any("NEEDS_VERIFICATION" in warning.value for warning in app.warning), (
        [error.value for error in app.error],
        app.session_state.get("test_api_calls"),
    )
    assert any(
        "Requirement, evidence and gap review" in header.value
        for header in app.markdown
    )
    assert not any(
        button.label == "Generate customised CV and cover letter"
        for button in app.button
    )

    fingerprint = app.session_state["manual_comparison"]["comparison_fingerprint"]
    app.checkbox(key=f"manual_approval_checked_{fingerprint}").check()
    app.button(key="manual_approve").click().run(timeout=10)
    assert any(
        "approval is recorded" in success.value.lower() for success in app.success
    )
    app.button(key="manual_generate").click().run(timeout=10)
    assert not app.exception, [(error.type, error.message) for error in app.exception]
    assert {button.label for button in app.download_button} == {
        "Download CV (DOCX)",
        "Download CV (PDF)",
        "Download Cover Letter (DOCX)",
        "Download Cover Letter (PDF)",
    }
    assert app.text_area(key="manual_cv_preview").value.startswith("# Candidate")
    assert len(app.session_state["test_api_calls"]) >= 7

    app.text_input(key="manual_job_title").set_value("Changed role title").run(
        timeout=10
    )
    assert not app.exception
    assert not any(
        button.label == "Generate customised CV and cover letter"
        for button in app.button
    )
    assert not app.download_button


def test_fetched_source_edits_are_clearly_marked_for_user_review():
    app = AppTest.from_file(
        str(PROJECT_ROOT / "tests/python/manual_application_fixture.py")
    ).run(timeout=10)
    app.text_input(key="manual_job_url").set_value("https://example.test/job")
    app.text_input(key="manual_job_title").set_value("Data Architect")
    app.text_input(key="manual_job_company").set_value("Example")
    app.button(key="manual_fetch_description").click().run(timeout=10)
    app.text_area(key="manual_job_description").set_value(
        app.text_area(key="manual_job_description").value + " Reviewed manually."
    ).run(timeout=10)
    app.button(key="manual_compare").click().run(timeout=10)

    assert not app.exception, [(error.type, error.message) for error in app.exception]
    assert any("user-edited" in warning.value for warning in app.warning)
    assert any("not attested by that fetch" in warning.value for warning in app.warning)


def test_manual_upload_reader_accepts_text_files_without_rewriting_content():
    class Uploaded:
        name = "role.txt"

        @staticmethod
        def getvalue():
            return b"Job title: Analyst\nBuild data models."

    assert (
        extract_uploaded_description(Uploaded())
        == "Job title: Analyst\nBuild data models."
    )


def test_manual_upload_reader_extracts_docx_paragraphs_and_tables():
    document = docx.Document()
    document.add_paragraph("Role summary")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Build data models"
    table.cell(0, 1).text = "Work with stakeholders"
    stream = BytesIO()
    document.save(stream)

    class Uploaded:
        name = "role.docx"

        @staticmethod
        def getvalue():
            return stream.getvalue()

    extracted = extract_uploaded_description(Uploaded())
    assert "Role summary" in extracted
    assert "Build data models\tWork with stakeholders" in extracted


def test_manual_upload_reader_rejects_oversized_expanded_docx():
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 5_000_001)

    class Uploaded:
        name = "large.docx"

        @staticmethod
        def getvalue():
            return stream.getvalue()

    with pytest.raises(ValueError, match="expands beyond"):
        extract_uploaded_description(Uploaded())


def test_download_bytes_validates_contract_hash_and_file_signature():
    content = b"%PDF-1.4\nvalid test file\n%%EOF\n"
    artifact = {"type": "CV", "format": "PDF", "filename": "cv.pdf"}

    def api_request(method, path, timeout):
        assert method == "GET"
        assert path.endswith("/documents/cv/pdf")
        return {
            "schema_version": "manual_application_v1",
            "ok": True,
            "filename": "download.pdf",
            "mime_type": "application/pdf",
            "sha256": hashlib.sha256(content).hexdigest(),
            "content_base64": base64.b64encode(content).decode("ascii"),
        }

    assert _download_bytes(
        api_request, "11111111-1111-4111-8111-111111111111", artifact
    ) == (content, "download.pdf")

    def tampered_response(method, path, timeout):
        result = api_request(method, path, timeout)
        result["sha256"] = "0" * 64
        return result

    with pytest.raises(ValueError, match="SHA-256"):
        _download_bytes(
            tampered_response, "11111111-1111-4111-8111-111111111111", artifact
        )


def test_python_rejects_api_payloads_that_do_not_match_versioned_zod_contracts():
    with pytest.raises(ValueError, match="contract violation"):
        _validate_payload(
            "ManualApprovalResponse",
            {
                "schema_version": "manual_application_v1",
                "ok": True,
                "approval_id": "not-a-uuid",
                "approved": True,
                "comparison_fingerprint": "a" * 64,
            },
        )
