"""User-supervised, single-job application preparation UI."""

import base64
import hashlib
import json
import re
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import docx
from jsonschema import Draft202012Validator, ValidationError
import streamlit as st


PROJECT_ROOT = Path(__file__).resolve().parent
SCHEMA_DIRECTORY = PROJECT_ROOT / "src" / "contracts" / "json"
_SCHEMA_CACHE = {}


def _validate_payload(schema_name, payload):
    schema = _SCHEMA_CACHE.get(schema_name)
    if schema is None:
        schema_path = SCHEMA_DIRECTORY / f"{schema_name}.schema.json"
        with schema_path.open("r", encoding="utf-8") as schema_file:
            schema = json.load(schema_file)
        Draft202012Validator.check_schema(schema)
        _SCHEMA_CACHE[schema_name] = schema
    try:
        Draft202012Validator(schema).validate(payload)
    except ValidationError as error:
        location = ".".join(str(part) for part in error.absolute_path) or "<root>"
        raise ValueError(f"API {schema_name} contract violation at {location}: {error.message}") from error
    return payload


def _validate_document_pair(artifacts):
    keys = {(item.get("type"), item.get("format")) for item in artifacts}
    expected = {("CV", "DOCX"), ("CV", "PDF"), ("COVER_LETTER", "DOCX"), ("COVER_LETTER", "PDF")}
    if len(artifacts) != 4 or keys != expected:
        raise ValueError("The API did not return one DOCX and one PDF for both documents.")


def extract_uploaded_description(uploaded_file):
    if uploaded_file is None:
        return ""
    name = (uploaded_file.name or "").lower()
    data = uploaded_file.getvalue()
    if len(data) > 1_000_000:
        raise ValueError("The uploaded source file is larger than the 1 MB manual import limit.")
    if name.endswith(".txt"):
        text = data.decode("utf-8-sig")
        if len(text) > 100_000:
            raise ValueError("The extracted description exceeds the 100,000 character limit.")
        return text
    if name.endswith(".docx"):
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 1_000 or sum(item.file_size for item in members) > 5_000_000:
                raise ValueError("The uploaded Word file expands beyond the 5 MB import limit.")
        document = docx.Document(BytesIO(data))
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            blocks.extend("\t".join(cell.text.strip() for cell in row.cells if cell.text.strip()) for row in table.rows)
        text = "\n".join(block for block in blocks if block.strip())
        if len(text) > 100_000:
            raise ValueError("The extracted description exceeds the 100,000 character limit.")
        return text
    raise ValueError("Upload a plain text (.txt) or Word (.docx) job description.")


def _download_bytes(api_request, draft_id, artifact):
    response = api_request(
        "GET",
        f"/api/v2/application-drafts/{draft_id}/documents/{artifact['type'].lower()}/{artifact['format'].lower()}",
        timeout=60,
    )
    _validate_payload("ManualDownloadResponse", response)
    if not response.get("ok") or not response.get("content_base64"):
        raise RuntimeError(response.get("error") or "The API did not return a document artifact.")
    content = base64.b64decode(response["content_base64"], validate=True)
    if hashlib.sha256(content).hexdigest() != response["sha256"]:
        raise ValueError("The downloaded private artifact failed its SHA-256 integrity check.")
    expected_mime = {
        "PDF": ("application/pdf", b"%PDF-"),
        "DOCX": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", b"PK\x03\x04"),
    }
    mime, signature = expected_mime[artifact["format"]]
    if response["mime_type"] != mime or not content.startswith(signature):
        raise ValueError("The downloaded file type does not match its validated artifact metadata.")
    filename = response["filename"]
    if not re.fullmatch(r"[A-Za-z0-9_. -]{1,160}", filename):
        raise ValueError("The private document filename is invalid.")
    return content, filename


def render_manual_application_flow(api_request):
    st.subheader("Prepare documents for one manually selected role")
    st.caption("Enter the posting URL even if you paste the description. This flow stages and compares only that role; it does not drain the ingestion or evaluation backlog.")

    candidate_profile_id = None
    try:
        candidate_response = api_request("GET", "/api/v2/application-drafts/candidate-profiles", timeout=30)
        _validate_payload("ManualCandidateProfilesResponse", candidate_response)
        profiles = candidate_response.get("profiles") or []
        if profiles:
            profile_by_id = {profile["candidate_profile_id"]: profile for profile in profiles}
            candidate_profile_id = st.selectbox(
                "Candidate profile",
                list(profile_by_id),
                format_func=lambda profile_id: f"{profile_by_id[profile_id]['display_name']} · {profile_id[:8]}",
                key="manual_candidate_profile",
            )
        else:
            st.error("No active candidate profile is available in this workspace. Activate a profile before comparing a role.")
    except Exception as error:
        st.error(f"Candidate profiles are unavailable: {error}")

    pending_url = st.session_state.pop("manual_pending_job_url", None)
    if pending_url:
        st.session_state["manual_job_url"] = pending_url
    job_url = st.text_input("Job URL", key="manual_job_url", placeholder="https://company.example/careers/role")
    title = st.text_input("Confirm job title", key="manual_job_title")
    company = st.text_input("Confirm company", key="manual_job_company")
    requisition_id = st.text_input("Requisition ID (optional)", key="manual_job_requisition_id")
    fetch_col, fetch_help_col = st.columns([1, 3])
    with fetch_col:
        fetch_clicked = st.button("Fetch description", key="manual_fetch_description")
    with fetch_help_col:
        st.caption("If fetching is blocked by the posting site, paste the description or upload a .txt/.docx file below.")
    fetch_completed = False
    if fetch_clicked:
        if not job_url.strip() or not title.strip() or not company.strip():
            st.error("Enter the URL and the title and company you expect the posting to contain.")
        else:
            try:
                fetched = api_request(
                    "POST", "/api/v2/application-drafts/fetch-description",
                    body={"url": job_url.strip(), "title": title.strip(), "company": company.strip(),
                          "requisition_id": requisition_id.strip() or None}, timeout=35,
                )
                _validate_payload("ManualFetchDescriptionResponse", fetched)
                if not fetched.get("ok") or not fetched.get("description"):
                    raise RuntimeError(fetched.get("error") or "No description was returned.")
                st.session_state["manual_job_description"] = fetched["description"]
                st.session_state["manual_pending_job_url"] = fetched.get("final_url") or job_url.strip()
                st.session_state["manual_job_description_source_hash"] = hashlib.sha256(fetched["description"].strip().encode("utf-8")).hexdigest()
                st.session_state["manual_job_source_kind"] = "FETCHED"
                st.session_state["manual_fetch_receipt_id"] = fetched["fetch_receipt_id"]
                source_evidence = fetched["source_evidence"]
                st.session_state["manual_fetch_identity"] = {
                    "urls": {source_evidence["requestedUrl"], source_evidence["resolvedUrl"], fetched["final_url"]},
                    "title": title.strip(),
                    "company": company.strip(),
                    "requisition_id": requisition_id.strip() or None,
                }
                st.session_state["manual_fetch_error"] = None
                fetch_completed = True
            except Exception as error:
                st.session_state["manual_fetch_error"] = str(error)
    if fetch_completed:
        st.rerun()
    if st.session_state.get("manual_fetch_error"):
        st.warning(f"Description fetch failed: {st.session_state['manual_fetch_error']} Paste or upload the description to continue.")

    uploaded_file = st.file_uploader("Upload description (.txt or .docx)", type=["txt", "docx"], key="manual_job_upload")
    if uploaded_file is not None:
        upload_marker = (uploaded_file.name, hashlib.sha256(uploaded_file.getvalue()).hexdigest())
        if st.session_state.get("manual_job_upload_marker") != upload_marker:
            try:
                st.session_state["manual_job_description"] = extract_uploaded_description(uploaded_file)
                st.session_state["manual_job_upload_marker"] = upload_marker
                st.session_state["manual_job_source_kind"] = "UPLOAD"
                st.session_state.pop("manual_fetch_receipt_id", None)
                st.session_state.pop("manual_fetch_identity", None)
            except Exception as error:
                st.error(f"Could not read uploaded description: {error}")
    description = st.text_area(
        "Confirm or edit the imported job description",
        key="manual_job_description",
        height=260,
        help="Check that the title, company and text belong to the exact posting you want to prepare for.",
    )

    input_signature = (job_url.strip(), title.strip(), company.strip(), description.strip(), candidate_profile_id)
    previous_signature = st.session_state.get("manual_preparation_input_signature")
    if previous_signature is not None and previous_signature != input_signature:
        st.session_state.pop("manual_comparison", None)
        st.session_state.pop("manual_approval", None)
        st.session_state.pop("manual_generated_documents", None)
        fetch_identity = st.session_state.get("manual_fetch_identity")
        if fetch_identity and (
            job_url.strip() not in fetch_identity["urls"]
            or title.strip().casefold() != fetch_identity["title"].casefold()
            or company.strip().casefold() != fetch_identity["company"].casefold()
            or (requisition_id.strip() or None) != fetch_identity["requisition_id"]
        ):
            st.session_state.pop("manual_fetch_receipt_id", None)
            st.session_state.pop("manual_fetch_identity", None)
            if st.session_state.get("manual_job_source_kind") == "FETCHED":
                st.session_state["manual_job_source_kind"] = "PASTED"
    st.session_state["manual_preparation_input_signature"] = input_signature

    compare_clicked = st.button("Compare with my profile", type="primary", key="manual_compare")
    comparison_completed = False
    if compare_clicked:
        if not candidate_profile_id:
            st.error("Select an active candidate profile before comparing this role.")
        elif not job_url.strip() or not title.strip() or not company.strip() or len(description.strip()) < 100:
            st.error("Provide the posting URL, confirm the title and company, and enter at least 100 characters of job description.")
        elif len(description) > 100_000:
            st.error("The confirmed job description exceeds the 100,000 character limit.")
        else:
            try:
                source_kind = st.session_state.get("manual_job_source_kind")
                if source_kind not in {"UPLOAD", "FETCHED"}:
                    source_kind = "PASTED"
                compare_body = {
                    "schema_version": "manual_application_v1",
                    "candidate_profile_id": candidate_profile_id,
                    "source_kind": source_kind,
                    "title": title.strip(), "company": company.strip(), "url": job_url.strip(),
                    "requisition_id": requisition_id.strip() or None,
                    "description": description.strip(), "location": "Unknown",
                }
                if source_kind == "FETCHED":
                    receipt_id = st.session_state.get("manual_fetch_receipt_id")
                    if not receipt_id:
                        raise ValueError("The verified fetch receipt is no longer available. Fetch the posting again or paste/upload its description.")
                    compare_body["fetch_receipt_id"] = receipt_id
                comparison = api_request(
                    "POST",
                    "/api/v2/application-drafts/compare",
                    body=compare_body,
                    timeout=120,
                )
                _validate_payload("ManualCompareResponse", comparison)
                if not comparison.get("ok"):
                    raise RuntimeError(comparison.get("error") or "Comparison failed.")
                st.session_state["manual_comparison"] = comparison
                st.session_state.pop("manual_approval", None)
                st.session_state.pop("manual_generated_documents", None)
                comparison_completed = True
            except Exception as error:
                st.error(f"Could not compare this role: {error}")
    if comparison_completed:
        st.rerun()

    comparison = st.session_state.get("manual_comparison")
    if not comparison:
        return
    st.markdown("### Selected posting")
    selected = comparison.get("selected_job") or {}
    st.text(f"{selected.get('title') or title} at {selected.get('company') or company}")
    if selected.get("url"):
        st.caption("Version-pinned posting URL")
        st.code(selected["url"], language=None)
    source = selected.get("source") or {}
    if source.get("source_kind") == "FETCHED":
        evidence = source.get("source_evidence") or {}
        st.caption(
            f"Verified source: {source.get('provider')} · retrieved {evidence.get('retrievedAt')} · "
            f"HTTP {evidence.get('httpStatus')} · response SHA-256 {str(evidence.get('responseHash', ''))[:12]}…"
        )
        if source.get("source_content_status") == "USER_EDITED":
            st.warning("You edited the fetched description. Its retrieval evidence is retained, but the staged text is marked user-edited and is not attested by that fetch.")
        else:
            st.info("The staged description matches the verified fetched content. The retrieval evidence is pinned to this posting observation.")
    else:
        st.caption("Source: user-pasted or uploaded description; no public retrieval receipt is attached.")

    payload = comparison.get("comparison") or {}
    for warning in payload.get("suitability_warnings", []):
        st.warning(warning)
    st.info(payload.get("evidence_notice") or "Review every potential evidence match and gap yourself.")
    rows = []
    for row in payload.get("rows", []):
        evidence = row.get("evidence") or []
        rows.append({
            "Requirement": row.get("requirement", ""),
            "Importance": row.get("importance", ""),
            "Potential profile evidence": "\n\n".join(item.get("statement", "") for item in evidence) or "No direct statement found",
            "Gap / review": row.get("gap") or "Review whether the potential evidence actually meets this requirement.",
        })
    st.markdown("### Requirement, evidence and gap review")
    if rows:
        st.dataframe(rows, use_container_width=True, hide_index=True)
    else:
        st.warning("No deterministic requirements were extracted. Review the description and confirm the gap manually before approval.")

    if comparison.get("document_contact_ready") is False:
        missing = ", ".join(comparison.get("missing_contact_fields", [])) or "required fields"
        st.error(f"The active private profile contact graph is missing: {missing}. Update those private profile contact facts before document preparation.")
        return

    draft_id = comparison.get("draft_id")
    fingerprint = comparison.get("comparison_fingerprint")
    approval_key = f"manual_approval_checked_{fingerprint}"
    st.checkbox("I screened this role and want to prepare documents", key=approval_key)
    if st.button("Approve document preparation", key="manual_approve"):
        if not st.session_state.get(approval_key):
            st.error("Check the explicit approval box after reviewing the job and evidence table.")
        else:
            try:
                result = api_request(
                    "POST",
                    f"/api/v2/application-drafts/{draft_id}/approve",
                    body={"approved": True, "comparison_fingerprint": fingerprint},
                    timeout=60,
                )
                _validate_payload("ManualApprovalResponse", result)
                if not result.get("ok"):
                    raise RuntimeError(result.get("error") or "Approval was not recorded.")
                st.session_state["manual_approval"] = result
                st.success("Your approval is recorded for this exact job version, profile version, requirement set and comparison.")
            except Exception as error:
                st.error(f"Could not record approval: {error}")

    approval = st.session_state.get("manual_approval")
    if not approval or not draft_id:
        return
    if st.button("Generate customised CV and cover letter", type="primary", key="manual_generate"):
        try:
            result = api_request(
                "POST",
                f"/api/v2/application-drafts/{draft_id}/documents",
                body={"approval_id": approval.get("approval_id")},
                timeout=120,
            )
            _validate_payload("ManualDocumentsResponse", result)
            _validate_document_pair(result.get("artifacts") or [])
            if not result.get("ok"):
                raise RuntimeError(result.get("error") or "Document generation failed.")
            st.session_state["manual_generated_documents"] = result
        except Exception as error:
            st.error(f"Could not generate the document pair: {error}")

    generated = st.session_state.get("manual_generated_documents")
    if not generated:
        return
    st.markdown("### Review before downloading")
    preview_col, letter_col = st.columns(2)
    with preview_col:
        st.text_area("Customised CV preview", value=(generated.get("preview") or {}).get("cv", ""), height=400, disabled=True, key="manual_cv_preview")
    with letter_col:
        st.text_area("Cover letter preview", value=(generated.get("preview") or {}).get("cover_letter", ""), height=400, disabled=True, key="manual_cover_letter_preview")
    st.caption("Review both previews and download only the private artifacts you intend to use. No application status is recorded.")
    for artifact in generated.get("artifacts", []):
        document_name = "CV" if artifact.get("type") == "CV" else "Cover Letter"
        label = f"Download {document_name} ({artifact.get('format', '')})"
        try:
            content, filename = _download_bytes(api_request, draft_id, artifact)
            st.download_button(label, content, file_name=filename, mime=artifact.get("mime_type", "application/octet-stream"),
                key=f"manual_download_{artifact.get('type')}_{artifact.get('format')}_{fingerprint}")
        except Exception as error:
            st.error(f"Private download is unavailable: {error}")
